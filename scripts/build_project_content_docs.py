#!/usr/bin/env python3
from __future__ import annotations

import json
import os
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "documents"
NODE = Path("/Users/reznikmd/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node")
NODE_MODULES = "/Users/reznikmd/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules"
EXPORTER = ROOT / "scripts" / "export_project_content.js"
STRUCTURE_PATH = OUT_DIR / "Структура_проекта_Код_жизни.docx"
CONTENT_PATH = OUT_DIR / "Полная_текстовка_проекта_Код_жизни.docx"

BLACK = "000000"
NAVY = "24364B"
PALE_BLUE = "EAF0F6"
PALE_GRAY = "F5F6F7"
GRID = "D9D9D9"


def load_data() -> dict:
    env = dict(os.environ)
    env["NODE_PATH"] = NODE_MODULES
    result = subprocess.run(
        [str(NODE), str(EXPORTER)],
        cwd=ROOT,
        env=env,
        check=True,
        stdout=subprocess.PIPE,
        text=True,
    )
    return json.loads(result.stdout)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRID)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5A94")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def font_run(run, name="Aptos", size=None, bold=None, color=BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def setup_document(title: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(12)

    for style_name, size, before, after in (
        ("Heading 1", 17, 16, 7),
        ("Heading 2", 14, 12, 5),
        ("Heading 3", 11.5, 9, 3),
        ("Heading 4", 10.5, 7, 2),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Compact List" not in styles:
        compact = styles.add_style("Compact List", WD_STYLE_TYPE.PARAGRAPH)
    else:
        compact = styles["Compact List"]
    compact.base_style = styles["Normal"]
    compact.paragraph_format.left_indent = Inches(0.22)
    compact.paragraph_format.first_line_indent = Inches(-0.16)
    compact.paragraph_format.space_after = Pt(1.5)
    compact.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.text = "Код жизни"
    font_run(header.runs[0], size=8.5, bold=True, color="5B6570")
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(section.footer.paragraphs[0])

    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(subtitle)
    font_run(r, size=12, color="3D4650")
    p.paragraph_format.space_after = Pt(18)

    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(5.2)
    labels = (("Проект", "Код жизни"), ("Состояние", "Фактическая структура и содержание на 20 сентября 2026 года"))
    for i, (label, value) in enumerate(labels):
        for cell in table.rows[i].cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(table.cell(i, 0), PALE_BLUE)
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
        for run in table.cell(i, 0).paragraphs[0].runs:
            font_run(run, size=9.5, bold=True)
        for run in table.cell(i, 1).paragraphs[0].runs:
            font_run(run, size=9.5)
    set_table_borders(table)
    doc.add_paragraph()
    return doc


def add_intro(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(10)


def add_compact_item(doc: Document, text: str, level=0) -> None:
    p = doc.add_paragraph(style="Compact List")
    p.paragraph_format.left_indent = Inches(0.22 + 0.24 * level)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.add_run("• ")
    p.add_run(text)


def add_content_item(doc: Document, text: str) -> None:
    if text == "":
        doc.add_paragraph()
        return
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("• ")
    p.add_run(str(text))


def add_document_break(doc: Document) -> None:
    doc.add_section(WD_SECTION_START.NEW_PAGE)


def build_structure(data: dict) -> Document:
    doc = setup_document(
        "Структура проекта Код жизни",
        "Краткая карта разделов приложения для визуальной проверки состава и иерархии",
    )
    add_intro(doc, "Документ показывает существующие блоки без раскрытия медицинских текстов. По нему удобно оценить полноту справочника, последовательность разделов и места для будущих дополнений.")

    refs = data["referenceCategories"]
    checks = data["checklistCategories"]
    campaigns = data["caseCampaigns"]
    summary = doc.add_table(rows=1, cols=3)
    summary.autofit = False
    widths = (2.3, 2.3, 2.3)
    values = (
        ("Справочник", f"{len(refs)} категорий"),
        ("Чек листы", f"{len(checks)} категорий"),
        ("Кейсы", f"{len(campaigns)} кампаний"),
    )
    for i, (label, value) in enumerate(values):
        cell = summary.cell(0, i)
        cell.width = Inches(widths[i])
        set_cell_margins(cell, top=160, bottom=160)
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label + "\n")
        font_run(r, size=10, bold=True, color="FFFFFF")
        r = p.add_run(value)
        font_run(r, size=9, color="FFFFFF")
    set_table_borders(summary)

    doc.add_heading("Справочник", level=1)
    for ci, category in enumerate(refs, 1):
        doc.add_heading(f"{ci}. {category['title']}", level=2)
        for si, section in enumerate(category.get("sections", []), 1):
            doc.add_heading(f"{ci}.{si} {section['title']}", level=3)
            for subsection in section.get("subsections", []):
                add_compact_item(doc, subsection["title"], 0)
                for article in subsection.get("articles", []):
                    add_compact_item(doc, article["title"], 1)
                for term in subsection.get("terms", []):
                    add_compact_item(doc, term["term"], 1)

    add_document_break(doc)
    doc.add_heading("Чек листы", level=1)
    for ci, category in enumerate(checks, 1):
        doc.add_heading(f"{ci}. {category['title']}", level=2)
        for checklist in category.get("checklists", []):
            add_compact_item(doc, checklist["title"])

    add_document_break(doc)
    doc.add_heading("Кейсы", level=1)
    for ci, campaign in enumerate(campaigns, 1):
        doc.add_heading(f"{ci}. {campaign['title']}", level=2)
        for scenario in campaign.get("scenarios", []):
            add_compact_item(doc, scenario["title"])
    return doc


FIELD_LABELS = {
    "goal": "Цель",
    "lookFor": "Что искать",
    "actions": "Действия",
    "resultCheck": "Проверка результата",
    "mistakes": "Типичные ошибки",
    "nextStep": "Следующий шаг",
    "context": "Контекст",
    "brief": "Вводная",
    "roles": "Роли",
    "equipment": "Оснащение",
    "timeline": "Хронология",
    "decisionOptions": "Варианты решений",
    "criticalErrors": "Критические ошибки",
    "successCriteria": "Критерии успеха",
    "legalBoundaries": "Правовые границы",
    "debriefQuestions": "Вопросы разбора",
}


def add_labeled_value(doc: Document, label: str, value) -> None:
    if value is None or value == "" or value == []:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    font_run(r, size=10.5, bold=True)
    if isinstance(value, list):
        for item in value:
            add_content_item(doc, str(item))
    else:
        p = doc.add_paragraph(str(value))
        p.paragraph_format.space_after = Pt(4)


def build_full_content(data: dict) -> Document:
    doc = setup_document(
        "Полная текстовка проекта Код жизни",
        "Редакционная выгрузка всего пользовательского содержания приложения для ручной проверки",
    )
    add_intro(doc, "Документ последовательно воспроизводит тексты справочника, чек листов и учебных кейсов. Служебные идентификаторы и программная разметка исключены; ссылки на источники и подписи изображений сохранены.")

    refs = data["referenceCategories"]
    checks = data["checklistCategories"]
    campaigns = data["caseCampaigns"]

    doc.add_heading("Часть 1 Справочник", level=1)
    for ci, category in enumerate(refs, 1):
        if ci > 1:
            add_document_break(doc)
        doc.add_heading(f"{ci}. {category['title']}", level=1)
        for si, section in enumerate(category.get("sections", []), 1):
            doc.add_heading(f"{ci}.{si} {section['title']}", level=2)
            if section.get("description"):
                p = doc.add_paragraph(section["description"])
                p.paragraph_format.space_after = Pt(7)
            for subi, subsection in enumerate(section.get("subsections", []), 1):
                doc.add_heading(subsection["title"], level=3)
                for article in subsection.get("articles", []):
                    doc.add_heading(article["title"], level=4)
                    if article.get("meta"):
                        p = doc.add_paragraph()
                        for mi, item in enumerate(article["meta"]):
                            if mi:
                                p.add_run("   ")
                            r = p.add_run(f"{item['label']}: ")
                            font_run(r, size=9, bold=True, color="5B6570")
                            r = p.add_run(str(item["value"]))
                            font_run(r, size=9, color="5B6570")
                    if article.get("image"):
                        image = article["image"]
                        p = doc.add_paragraph()
                        r = p.add_run("Изображение: ")
                        font_run(r, size=9, bold=True, color="5B6570")
                        r = p.add_run(image.get("alt", ""))
                        font_run(r, size=9, color="5B6570")
                        if image.get("caption"):
                            r = p.add_run(f". {image['caption']}")
                            font_run(r, size=9, color="5B6570")
                    for item in article.get("content", []):
                        add_content_item(doc, item)
                    for source in article.get("links", []):
                        p = doc.add_paragraph()
                        p.paragraph_format.left_indent = Inches(0.2)
                        p.paragraph_format.space_after = Pt(2)
                        add_hyperlink(p, source["label"], source["href"])
                for term in subsection.get("terms", []):
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(3)
                    r = p.add_run(term["term"] + ". ")
                    font_run(r, size=10.5, bold=True)
                    p.add_run(term["definition"])

    add_document_break(doc)
    doc.add_heading("Часть 2 Чек листы", level=1)
    for ci, category in enumerate(checks, 1):
        if ci > 1:
            add_document_break(doc)
        doc.add_heading(f"{ci}. {category['title']}", level=1)
        for checklist in category.get("checklists", []):
            doc.add_heading(checklist["title"], level=2)
            for key, value in checklist.items():
                if key in ("id", "title"):
                    continue
                add_labeled_value(doc, FIELD_LABELS.get(key, key), value)

    add_document_break(doc)
    doc.add_heading("Часть 3 Кейсы", level=1)
    for ci, campaign in enumerate(campaigns, 1):
        if ci > 1:
            add_document_break(doc)
        doc.add_heading(f"{ci}. {campaign['title']}", level=1)
        for scenario in campaign.get("scenarios", []):
            doc.add_heading(scenario["title"], level=2)
            for key, value in scenario.items():
                if key in ("id", "title"):
                    continue
                add_labeled_value(doc, FIELD_LABELS.get(key, key), value)
    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    data = load_data()
    structure = build_structure(data)
    structure.save(STRUCTURE_PATH)
    full = build_full_content(data)
    full.save(CONTENT_PATH)
    print(STRUCTURE_PATH)
    print(CONTENT_PATH)


if __name__ == "__main__":
    main()
